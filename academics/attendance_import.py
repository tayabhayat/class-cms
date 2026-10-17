
from datetime import date, datetime

from openpyxl import load_workbook

STATUS_MAP = {
    "P": "present",
    "P(LATE)": "present",
    "P (LATE)": "present",
    "PRESENT": "present",
    "A": "absent",
    "ABSENT": "absent",
    "L": "leave",
    "LEAVE": "leave",
}


class AttendanceRecordList(list):
    def __init__(self, items=None, skipped_count=0):
        super().__init__(items or [])
        self.skipped_count = skipped_count


def parse_attendance_excel(uploaded_file):
    workbook = load_workbook(
        uploaded_file,
        read_only=True,
        data_only=True,
    )

    try:
        sheet = workbook.active

        headers = [
            sheet.cell(row=2, column=col).value
            for col in range(1, sheet.max_column + 1)
        ]

        id_col = next(
            (i + 1 for i, value in enumerate(headers)
             if str(value or "").strip().lower() == "student id"),
            None,
        )

        name_col = next(
            (i + 1 for i, value in enumerate(headers)
             if str(value or "").strip().lower() == "student name"),
            None,
        )

        if not id_col or not name_col:
            raise ValueError(
                "Could not find Student ID and Student Name in row 2."
            )

        date_columns = {}
        seen_dates = set()
        errors = []
        max_header_col = max(id_col, name_col)

        for col in range(max_header_col + 1, sheet.max_column + 1):
            value = sheet.cell(row=2, column=col).value

            if value is None:
                continue

            parsed_date = None
            if isinstance(value, datetime):
                parsed_date = value.date()
            elif isinstance(value, date):
                parsed_date = value
            elif isinstance(value, str):
                val_str = value.strip()
                if not val_str:
                    continue
                for fmt in ("%Y-%m-%d", "%d-%m-%Y", "%m/%d/%Y", "%d/%m/%Y", "%Y/%m/%d"):
                    try:
                        parsed_date = datetime.strptime(val_str, fmt).date()
                        break
                    except ValueError:
                        pass
                if parsed_date is None:
                    errors.append(
                        f"Row 2, column {col}: invalid date header {value!r}."
                    )
                    continue
            else:
                errors.append(
                    f"Row 2, column {col}: invalid date header {value!r}."
                )
                continue

            if parsed_date in seen_dates:
                errors.append(
                    f"Row 2, column {col}: duplicate attendance date {parsed_date.isoformat()}."
                )
                continue

            seen_dates.add(parsed_date)
            date_columns[col] = parsed_date

        if not date_columns:
            raise ValueError("No attendance dates were found in row 2.")

        records = []
        seen_ids = set()
        skipped_count = 0

        for row_num in range(3, sheet.max_row + 1):
            raw_id = sheet.cell(row=row_num, column=id_col).value
            raw_name = sheet.cell(row=row_num, column=name_col).value

            if raw_id is None or not str(raw_id).strip():
                continue

            student_id = str(raw_id).strip().upper()

            if student_id in seen_ids:
                errors.append(
                    f"Row {row_num}: duplicate student ID {student_id}."
                )
                continue

            seen_ids.add(student_id)

            for col, attendance_date in date_columns.items():
                raw_status = sheet.cell(row=row_num, column=col).value

                # Blank cells must never erase existing attendance.
                if raw_status is None or not str(raw_status).strip():
                    skipped_count += 1
                    continue

                status_key = str(raw_status).strip().upper()
                status = STATUS_MAP.get(status_key)

                if status is None:
                    errors.append(
                        f"Row {row_num}, {attendance_date}: "
                        f"unknown status {raw_status!r}."
                    )
                    continue

                records.append({
                    "student_id": student_id,
                    "student_name": str(raw_name or "").strip(),
                    "date": attendance_date.isoformat(),
                    "status": status,
                })

        return AttendanceRecordList(records, skipped_count=skipped_count), errors

    finally:
        workbook.close()
