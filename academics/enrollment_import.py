
from openpyxl import load_workbook


def parse_enrollment_excel(uploaded_file):
    workbook = load_workbook(
        uploaded_file,
        read_only=True,
        data_only=True,
    )

    try:
        sheet = workbook.active

        headers = [
            sheet.cell(row=1, column=col).value
            for col in range(1, sheet.max_column + 1)
        ]

        id_col = next(
            (
                i + 1
                for i, value in enumerate(headers)
                if str(value or "").strip().lower() == "student id"
            ),
            None,
        )

        if id_col is None:
            raise ValueError(
                "Could not find a Student ID column in row 1."
            )

        student_ids = []
        errors = []
        seen_ids = set()

        for row_num in range(2, sheet.max_row + 1):
            raw_id = sheet.cell(row=row_num, column=id_col).value

            if raw_id is None or not str(raw_id).strip():
                continue

            student_id = str(raw_id).strip().upper()

            if student_id in seen_ids:
                errors.append(
                    f"Row {row_num}: duplicate student ID {student_id}."
                )
                continue

            seen_ids.add(student_id)
            student_ids.append(student_id)

        if not student_ids:
            errors.append("No student IDs were found in the spreadsheet.")

        return student_ids, errors

    finally:
        workbook.close()
