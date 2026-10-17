
import csv
import secrets
from pathlib import Path

from django.contrib.auth import get_user_model
from django.core.management.base import BaseCommand, CommandError
from django.db import transaction
from openpyxl import load_workbook

User = get_user_model()


class Command(BaseCommand):
    help = "Import students from an Excel roster."

    def add_arguments(self, parser):
        parser.add_argument("--file", required=True)
        parser.add_argument("--limit", type=int, default=None)
        parser.add_argument(
            "--credentials-out",
            default="student_initial_credentials.csv",
        )

    def handle(self, *args, **options):
        file_path = Path(options["file"]).expanduser()
        limit = options["limit"]
        credentials_path = Path(
            options["credentials_out"]
        ).expanduser()

        if not file_path.is_file():
            raise CommandError(f"File not found: {file_path}")

        if limit is not None and limit < 1:
            raise CommandError("--limit must be at least 1.")

        workbook = load_workbook(
            file_path,
            read_only=True,
            data_only=True,
        )

        try:
            sheet = workbook["Students"]
            rows = sheet.iter_rows(values_only=True)

            headers = next(rows, None)
            if not headers:
                raise CommandError("The worksheet is empty.")

            header_map = {
                str(value).strip().lower(): index
                for index, value in enumerate(headers)
                if value is not None
            }

            required = ["student id", "student name", "email"]
            missing = [
                name for name in required
                if name not in header_map
            ]

            if missing:
                raise CommandError(
                    f"Missing required columns: {', '.join(missing)}"
                )

            students = []
            seen_ids = set()
            seen_emails = set()

            for row_number, row in enumerate(rows, start=2):
                def value(column):
                    index = header_map[column]
                    return (
                        str(row[index]).strip()
                        if index < len(row) and row[index] is not None
                        else ""
                    )

                student_id = value("student id").upper()
                name = value("student name")
                email = value("email").lower()

                if not student_id and not name and not email:
                    continue

                if not student_id or not name or not email:
                    raise CommandError(
                        f"Row {row_number}: ID, name and email "
                        "are all required."
                    )

                if len(student_id) > 20:
                    raise CommandError(
                        f"Row {row_number}: student ID is too long."
                    )

                if len(email) > 254 or "@" not in email:
                    raise CommandError(
                        f"Row {row_number}: invalid email address."
                    )

                if student_id in seen_ids:
                    raise CommandError(
                        f"Row {row_number}: duplicate ID {student_id}."
                    )

                if email in seen_emails:
                    raise CommandError(
                        f"Row {row_number}: duplicate email {email}."
                    )

                seen_ids.add(student_id)
                seen_emails.add(email)
                students.append({
                    "student_id": student_id,
                    "name": name,
                    "email": email,
                    "row": row_number,
                })

            if limit is not None:
                students = students[:limit]

        finally:
            workbook.close()

        if not students:
            raise CommandError("No student records found.")

        created = 0
        updated = 0
        credentials = []

        try:
            with transaction.atomic():
                for record in students:
                    student_id = record["student_id"]
                    email = record["email"]
                    name = record["name"]

                    first_name, _, last_name = name.partition(" ")

                    by_id = User.objects.filter(
                        student_id=student_id
                    ).first()

                    by_email = User.objects.filter(
                        email__iexact=email
                    ).first()

                    if by_id and by_id.role != User.Role.STUDENT:
                        raise CommandError(
                            f"Row {record['row']}: {student_id} "
                            "belongs to a non-student account."
                        )

                    if by_email and by_email != by_id:
                        if (
                            by_email.role != User.Role.STUDENT
                            or by_email.student_id not in (None, student_id)
                        ):
                            raise CommandError(
                                f"Row {record['row']}: email {email} "
                                "belongs to another account."
                            )

                    user = by_id or by_email

                    if user:
                        if (
                            user.student_id
                            and user.student_id != student_id
                        ):
                            raise CommandError(
                                f"Row {record['row']}: email is already "
                                "linked to a different student ID."
                            )

                        user.student_id = student_id
                        user.email = email
                        user.first_name = first_name
                        user.last_name = last_name
                        user.role = User.Role.STUDENT
                        user.save()
                        updated += 1
                    else:
                        temporary_password = secrets.token_urlsafe(12)

                        user = User.objects.create_user(
                            email=email,
                            password=temporary_password,
                            student_id=student_id,
                            first_name=first_name,
                            last_name=last_name,
                            role=User.Role.STUDENT,
                            is_active=True,
                        )

                        credentials.append({
                            "student_id": student_id,
                            "name": name,
                            "email": email,
                            "temporary_password": temporary_password,
                        })
                        created += 1

        except Exception as exc:
            if isinstance(exc, CommandError):
                raise
            raise CommandError(
                f"Import failed; database changes were rolled back: {exc}"
            ) from exc

        if credentials:
            credentials_path.parent.mkdir(
                parents=True,
                exist_ok=True,
            )
            write_header = not credentials_path.exists()

            with credentials_path.open(
                "a",
                newline="",
                encoding="utf-8",
            ) as output:
                writer = csv.DictWriter(
                    output,
                    fieldnames=[
                        "student_id",
                        "name",
                        "email",
                        "temporary_password",
                    ],
                )

                if write_header:
                    writer.writeheader()

                writer.writerows(credentials)

        self.stdout.write(self.style.SUCCESS(
            f"Import complete: {created} created, "
            f"{updated} updated, {len(students)} processed."
        ))

        if credentials:
            self.stdout.write(
                f"New-account credentials saved to: {credentials_path}"
            )
            self.stdout.write(self.style.WARNING(
                "Protect this file. It contains plaintext temporary "
                "passwords. Distribute credentials privately and "
                "delete the file securely when no longer needed."
            ))
